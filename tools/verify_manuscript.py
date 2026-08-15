#!/usr/bin/env python3
"""Mechanical DOCX checks for the WTCR Kindle manuscript."""
import argparse, sys
from docx import Document

def main():
    p=argparse.ArgumentParser(); p.add_argument('docx'); a=p.parse_args(); d=Document(a.docx)
    chapters=[x.text for x in d.paragraphs if x.style and x.style.name.startswith('Heading 1') and x.text.upper().startswith('CHAPTER')]
    scenes=sum(1 for x in d.paragraphs if x.text.strip()=='* * *')
    tabs=sum(x.text.count('\t') for x in d.paragraphs)
    ok=len(chapters)==46 and scenes==331 and tabs==0
    print(f'chapters={len(chapters)} scene_breaks={scenes} tabs={tabs} result={"PASS" if ok else "FAIL"}')
    sys.exit(0 if ok else 1)
if __name__=='__main__': main()
